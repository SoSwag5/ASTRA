"""Pre-change local backup and compact baseline inventory; no private text in stdout."""
import hashlib,json,sqlite3,shutil,sys
from pathlib import Path
from datetime import datetime
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from backend.models import engine,DATA
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
destination=DATA/'backups'/('campaign_baseline_'+datetime.now().strftime('%Y%m%d_%H%M%S'))
destination.mkdir(parents=True)
with sqlite3.connect(engine.url.database) as source,sqlite3.connect(destination/'database.db') as target:
    source.backup(target)
    source.row_factory=sqlite3.Row
    records=[dict(r) for r in source.execute('SELECT a.*,j.company,j.title,j.notes,j.date_found FROM applications a JOIN jobs j ON a.job_id=j.id ORDER BY a.id')]
    sources=[dict(r) for r in source.execute('SELECT * FROM job_sources')]
    job_count=source.execute('SELECT COUNT(*) FROM jobs').fetchone()[0]
raw=json.dumps(records,sort_keys=True,ensure_ascii=False)
(destination/'applications.json').write_text(raw,encoding='utf-8')
books=[]
for file in [DATA/'tracker.xlsx',root.parent/'Ayham_Job_Application_Tracker.xlsx']:
    if not file.exists():continue
    shutil.copy2(file,destination/file.name)
    book=load_workbook(file,read_only=False,data_only=False)
    books.append({'path':str(file),'sheets':[{'name':s.title,'rows':s.max_row,'columns':s.max_column,'tables':list(s.tables),'charts':len(s._charts),'populated_rows':sum(any(c.value is not None for c in row) for row in s)} for s in book]})
    book.close()
hashes={}
for folder in ['backend','frontend/src','scripts','tests']:
    for file in (root/folder).rglob('*'):
        if file.is_file() and '__pycache__' not in file.parts:
            relative=file.relative_to(root);hashes[str(relative)]=hashlib.sha256(file.read_bytes()).hexdigest()
            copy=destination/'source'/relative;copy.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(file,copy)
report={'backup':str(destination),'applications':len(records),'application_ids':[r['id'] for r in records],'preservation_sha256':hashlib.sha256(raw.encode()).hexdigest(),'jobs':job_count,'sources':sources,'workbooks':books,'source_hashes':hashes}
(destination/'inventory.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='source_hashes'},indent=2))
