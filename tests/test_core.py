import pytest,copy
from types import SimpleNamespace as NS
from sqlalchemy import create_engine,select
from sqlalchemy.orm import Session as OrmSession
from backend.models import *
from backend.policy import *
from backend.services import score,answer_for,add_job,generate_cv,sync_tracker,import_tracker,import_cv,set_status

@pytest.fixture
def db(tmp_path,monkeypatch):
    import backend.services as svc
    monkeypatch.setattr(svc,'DATA',tmp_path)
    e=create_engine('sqlite://'); Base.metadata.create_all(e)
    with OrmSession(e,expire_on_commit=False) as db:
        db.add(Settings(id=1,value=DEFAULTS)); db.commit(); yield db
def profile(): return {'raw_text':'Computer Science Python SIEM SOC log analysis Nmap Linux','name':'Candidate Example','email':'test@example.com','phone':'UNKNOWN','location':'Abu Dhabi'}
def job(**kw): return NS(title='Junior SOC Analyst',description='SOC SIEM log analysis Python. Entry level.',company='Example',location='Abu Dhabi',remote_status='On-site',job_url='https://jobs.example.com/123',match_score=90,recommendation='APPLY',duplicate_of=None,source='Company',**kw)
def test_duplicates():
    assert duplicate({'job_url':'https://x.com/jobs/1?utm_source=li'},{'job_url':'https://x.com/jobs/1'})
    assert duplicate({'company':'ACME','title':'SOC Analyst','location':'Dubai'},{'company':'Acme','title':'SOC Analyst','location':'Dubai'})
    assert not duplicate({'company':'Acme','title':'SOC Analyst','location':'Dubai'},{'company':'Other','title':'SOC Analyst','location':'Dubai'})
def test_duplicate_import(db):
    j,_=add_job(db,{'company':'A','title':'SOC','job_url':'https://linkedin.com/jobs/1'})
    j2,d=add_job(db,{'company':'A','title':'SOC','job_url':'https://example.com/jobs/1'})
    assert j.id==j2.id and d and j.job_url.startswith('https://example.com')
def test_scoring():
    result=score(job(),profile(),DEFAULTS); assert result['score']>=80; assert result['recommendation']=='HIGH_PRIORITY'
@pytest.mark.parametrize('requirement',['UAE National only','Active security clearance required','Arabic mandatory','CCNA required','5 years minimum required'])
def test_hard_requirement(requirement):
    j=job(); j.description=requirement; r=score(j,profile(),DEFAULTS); assert r['needs_review'] or r['hard_blockers']; assert r['recommendation']!='APPLY'
def test_question_memory(db):
    db.add(ApprovedAnswer(question='Do you need sponsorship?',normalized_question=norm('Do you need sponsorship?'),answer='Yes',approved=True));db.flush()
    assert answer_for(db,'Do you need sponsorship?',profile())=='Yes'
    assert answer_for(db,'Are you authorized to work?',profile()) is None
    assert answer_for(db,'Email',profile())=='test@example.com'
    assert answer_for(db,'Phone',profile()) is None
    assert answer_for(db,'What is your gender?',profile()) is None
@pytest.mark.parametrize('text',['CAPTCHA','Cloudflare verification','Two-factor authentication','Enter verification code','Sign in to continue'])
def test_challenges(text): assert challenge(text)
def test_transitions():
    transition('READY_TO_APPLY','APPLIED'); transition('APPLIED','INTERVIEW')
    with pytest.raises(ValueError): transition('APPLIED','READY_TO_APPLY')
    with pytest.raises(ValueError): transition('FOUND','NONSENSE')
def gate():
    return job(),NS(submission_attempted=False,status='READY_TO_APPLY',attempts=0),{**DEFAULTS,'autopilot':'AUTO_ALLOWED'},NS(enabled=True,auto_submit=True,permitted=True,domain='jobs.example.com')
def test_auto_policy():
    j,a,c,s=gate(); assert submission_gate(j,a,c,s,0,True)  # No site setting can enable external submission.
    assert submission_gate(j,a,c,None,0,True)
    assert submission_gate(j,a,c,s,10,True)
    assert submission_gate(j,a,c,s,0,True,True)
    a.attempts=3; assert submission_gate(j,a,c,s,0,True)
    a.attempts=0;a.submission_attempted=True; assert submission_gate(j,a,c,s,0,True)
def test_linkedin_manual():
    j,a,c,s=gate();j.job_url='https://www.linkedin.com/jobs/view/123';j.source='LinkedIn';s.domain='www.linkedin.com'
    assert submission_gate(j,a,c,s,0,True)
    with pytest.raises(ValueError): validate_url(j.job_url)
def test_ssrf():
    with pytest.raises(ValueError): validate_url('file:///etc/passwd')
    with pytest.raises(ValueError): validate_url('http://127.0.0.1')
def test_cv_preserves_facts(db,tmp_path):
    p=CandidateProfile(name='Candidate Example',email='test@example.com',phone='UNKNOWN',location='Abu Dhabi',summary='Security graduate.',raw_text='Security graduate.\nPython\nAcademic project only.',confirmed=True);db.add(p);db.flush()
    db.add(Skill(candidate_id=p.id,text='Python'));db.add(Project(candidate_id=p.id,text='Academic project only.'));db.flush()
    j,_=add_job(db,{'company':'X','title':'SOC','description':'CISSP Splunk Python'});r=generate_cv(db,j)
    assert 'Splunk' not in r.tailored and 'CISSP' not in r.tailored and 'Academic project only.' in r.tailored
    assert (tmp_path/r.pdf_path).exists() and (tmp_path/r.docx_path).exists()
def test_excel_roundtrip(db,tmp_path):
    j,_=add_job(db,{'company':'Example','title':'SOC','job_url':'https://example.com/job'});db.commit();sync_tracker(db)
    from openpyxl import load_workbook
    w=load_workbook(tmp_path/'tracker.xlsx');w.create_sheet('Keep me')['A1']='=1+2';w['Applications']['C2'].font=__import__('openpyxl').styles.Font(bold=True);w.save(tmp_path/'tracker.xlsx');w.close()
    sync_tracker(db);w=load_workbook(tmp_path/'tracker.xlsx');assert w['Keep me']['A1'].value=='=1+2';w.close()
    result=import_tracker(db,tmp_path/'tracker.xlsx');assert result['duplicates']==1;assert list((tmp_path/'backups').glob('*.xlsx'))
def test_followups(db):
    j,_=add_job(db,{'company':'A','title':'SOC'});set_status(db,j,'APPLIED');db.flush(); assert db.scalar(select(FollowUp));set_status(db,j,'APPLIED');db.flush();assert db.query(FollowUp).count()==1
