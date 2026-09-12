import copy,os,subprocess,sys
from datetime import datetime,timedelta,timezone
from pathlib import Path
import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session as OrmSession
from backend.models import Base,Settings,DEFAULTS,Job,Application,ApplicationEvent,FollowUp,WorkbookSync
from backend.campaign import analytics,application_rows,record_event
from backend.workbook import export_workbook

@pytest.fixture
def db():
    engine=create_engine('sqlite://')
    Base.metadata.create_all(engine)
    with OrmSession(engine,expire_on_commit=False) as db:
        db.add(Settings(id=1,value=DEFAULTS));db.commit();yield db

def make_rows(db,n=10):
    clock=datetime(2026,9,11,tzinfo=timezone.utc)
    for i in range(n):
        j=Job(company='Synthetic '+str(i),title='SOC Analyst',location='Dubai',source='Manual',notes='مرحبا\nUnicode José')
        db.add(j);db.flush()
        a=Application(job_id=j.id,status='APPLIED',applied_date=(clock-timedelta(days=30)).isoformat(),tracking={'stage':'APPLIED','cv_version':'SOC Focus','fit_at_application':'STRONG'})
        db.add(a);db.flush()
        if i<4:
            record_event(db,a,'MEANINGFUL_RESPONSE',when=(clock-timedelta(days=20)).isoformat())
            record_event(db,a,'MEANINGFUL_RESPONSE',when=(clock-timedelta(days=19)).isoformat())
        if i<2:record_event(db,a,'INTERVIEW',when=(clock-timedelta(days=15)).isoformat())
        if i<1:record_event(db,a,'OFFER',when=(clock-timedelta(days=10)).isoformat())
    db.commit();return clock

def test_metrics_exact_and_distinct(db):
    clock=make_rows(db)
    result=analytics(application_rows(db),14,clock)
    assert result['total']==10 and result['responses']==4 and result['interviews']==2 and result['offers']==1
    assert result['response_rate']==.4 and result['interview_rate']==.2 and result['offer_rate']==.1
    assert result['median_response_days']==10 and result['response_sample']==4

def test_metrics_empty_recent_and_receipts(db):
    assert analytics([])['response_rate'] is None
    clock=make_rows(db,1);a=db.query(Application).first();a.applied_date=(clock-timedelta(days=1)).isoformat()
    db.query(ApplicationEvent).delete();db.flush();record_event(db,a,'AUTOMATED_CONFIRMATION',when=clock.isoformat());db.commit()
    result=analytics(application_rows(db),14,clock)
    assert result['response_rate'] is None and result['responses']==0 and result['mature_count']==0
    a.status='WITHDRAWN';a.tracking={'stage':'WITHDRAWN'};db.commit()
    assert analytics(application_rows(db),14,clock)['active']==0

def test_historical_fit_does_not_change(db):
    make_rows(db,1);j=db.query(Job).first();j.match_score=0;j.analysis={'hard_blockers':['new rule']};db.commit()
    assert application_rows(db)[0]['tracking']['fit_at_application']=='STRONG'

def test_workbook_idempotence_unicode_and_legacy(db,tmp_path):
    from openpyxl import Workbook,load_workbook
    make_rows(db,5)
    target=tmp_path/'tracker.xlsx';w=Workbook();w.active.title='Applications';w.active.append(['ID','Company','Job Title']);w.active.append([99,'Legacy','Role']);w.create_sheet('Keep me')['A1']='=1+2';w.save(target);w.close()
    export_workbook(db,target);export_workbook(db,target)
    w=load_workbook(target);assert w['Applications'].max_row==6
    assert w['Discovered Jobs']['B2'].value=='Legacy' and w['Keep me']['A1'].value=='=1+2'
    assert 'مرحبا' in w['Applications']['U2'].value and len(w['Applications'].tables)==1
    assert len(w['Dashboard']._charts)==4
    w.close();assert len(list((tmp_path/'backups').glob('*.xlsx')))==2

def test_workbook_failed_replace_preserves_previous(db,tmp_path,monkeypatch):
    import backend.workbook as workbook
    make_rows(db,1);target=tmp_path/'tracker.xlsx';export_workbook(db,target);original=target.read_bytes()
    def fail(*args):raise PermissionError('Synthetic Excel lock')
    monkeypatch.setattr(workbook.os,'replace',fail)
    with pytest.raises(PermissionError):export_workbook(db,target)
    assert target.read_bytes()==original and not target.with_name('tracker.tmp.xlsx').exists()

def test_isolated_campaign_flow_and_preservation(tmp_path):
    script=r'''
from fastapi.testclient import TestClient
from sqlalchemy import select
from backend.main import app
from backend.models import *
from backend.campaign import application_rows
from backend.source_catalog import seed_catalog
from backend.workbook import retry_sync
from backend.reliability import backup_database
import json,sqlite3
with TestClient(app) as c:
    with Session.begin() as db:
        cfg=db.get(Settings,1);cfg.value={**cfg.value,'campaign_workbook':True,'discovery_enabled':False}
        seed_catalog(db)
        for i in range(5):
            j=Job(company='Old '+str(i),title='Analyst',notes='Original notes',date_found='2025-01-01');db.add(j);db.flush()
            db.add(Application(job_id=j.id,status='APPLIED',applied_date='2025-01-02',confirmation='original'))
    with Session() as db:
        before=[serialize(a) for a in db.scalars(select(Application))]
    initialize();initialize()
    with Session() as db:assert before==[serialize(a) for a in db.scalars(select(Application))]
    portals=c.get('/api/campaign/portals').json()
    assert len(portals['daily'])==8 and all(s['adapter']=='manual' and not s['enabled'] for s in portals['sources'])
    source=portals['daily'][0]
    assert c.post('/api/campaign/portals/'+str(source['id'])+'/checked',json={}).status_code==200
    for i,source in enumerate(['LinkedIn','Indeed','Bayt','GulfTalent','Naukrigulf','Dubai Careers','Lever','Ashby']):
        job=c.post('/api/jobs',json={'company':'New '+str(i),'title':'SOC Analyst','source':source,'location':['Dubai','Abu Dhabi','Sharjah'][i%3],'job_url':'https://example.com/job/'+str(i),'description':'Local synthetic SOC opportunity','notes':'العربية\nUnicode José'}).json()
        response=c.post('/api/campaign/jobs/'+str(job['id'])+'/track',json={'stage':'SHORTLISTED'})
        assert response.status_code==200,response.text
        response=c.post('/api/campaign/jobs/'+str(job['id'])+'/track',json={'stage':'APPLIED','cv_version':'SOC Focus','date':'2026-08-01'})
        assert response.status_code==200,response.text
        duplicate=c.post('/api/jobs',json={'company':'New '+str(i),'title':'SOC Analyst','source':'Other','location':job['location'],'job_url':job['job_url'],'notes':'Second occurrence'}).json()
        assert duplicate['duplicate'] and duplicate['id']==job['id']
        for stage in ['SCREENING','ASSESSMENT','INTERVIEW','OFFER','REJECTED']:
            response=c.post('/api/campaign/jobs/'+str(job['id'])+'/track',json={'stage':stage,'date':'2026-08-15'})
            assert response.status_code==200,response.text
        assert c.post('/api/campaign/jobs/'+str(job['id'])+'/track',json={'event_type':'MEANINGFUL_RESPONSE','date':'2026-08-05'}).status_code==200
    result=c.get('/api/campaign').json()
    assert result['metrics']['total']==13 and result['metrics']['interviews']==8 and result['metrics']['offers']==8
    retry_sync()
    assert (DATA/'tracker.xlsx').exists()
    backup=backup_database(force=True)
    with sqlite3.connect(DATA/'backups'/backup) as check:assert check.execute('SELECT COUNT(*) FROM applications').fetchone()[0]==13
    with Session() as db:assert before==[serialize(a) for a in db.scalars(select(Application).where(Application.id<=5))]
with TestClient(app) as c:assert c.get('/api/campaign').json()['metrics']['total']==13
'''
    data=tmp_path/'data';env={**os.environ,'HUNTER_DATA_DIR':str(data),'DATABASE_URL':f'sqlite:///{data/"test.db"}','APP_TOKEN':''}
    result=subprocess.run([sys.executable,'-c',script],env=env,text=True,capture_output=True,timeout=90)
    assert result.returncode==0,result.stdout+result.stderr
