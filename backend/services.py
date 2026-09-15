import re, json, hashlib, difflib, shutil, os, tempfile
from pathlib import Path
from datetime import datetime, timedelta, timezone
from copy import copy
from xml.sax.saxutils import escape
from sqlalchemy import select
from pypdf import PdfReader
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_LEFT
from docx import Document
from openpyxl import load_workbook, Workbook
from .models import *
from .policy import *
from .discovery import discovery_reason
from . import normalization, deduplication
from .normalization import JobObservationInput, canonical_view

def experience_years(description):
    from .recall import experience
    return experience(description)['minimum']

FACT_MODELS={'EXPERIENCE':Employment,'PERSONAL PROJECTS':Project,'EDUCATION':Education,'CERTIFICATIONS AND COURSES':Certification}
def import_cv(db,path):
    from .document_security import extract_pdf
    text=extract_pdf(path)
    if len(text)<100: raise ValueError('No readable CV text. Upload a text-based PDF.')
    profile=db.scalar(select(CandidateProfile))
    if not profile: profile=CandidateProfile(); db.add(profile); db.flush()
    lines=[x.strip() for x in text.splitlines() if x.strip()]
    profile.name=lines[0]
    email=re.search(r'[\w.+-]+@[\w.-]+\.[a-zA-Z]+',text)
    phone=re.search(r'\+\d[\d ()-]{7,}',text)
    profile.email=email.group() if email else 'UNKNOWN'
    profile.phone=phone.group().strip() if phone else 'UNKNOWN'
    profile.location=lines[1].split('|')[0].strip() if '|' in lines[1] else 'UNKNOWN'
    profile.raw_text=text; profile.confirmed=False
    profile.declarations={**(profile.declarations or {}),'extraction_state':'EXTRACTED — NEEDS CONFIRMATION','cv_language':'UNKNOWN'}
    for model in [Skill,*FACT_MODELS.values()]:
        for row in db.scalars(select(model).where(model.candidate_id==profile.id)): db.delete(row)
    headings=['PROFILE','TECHNICAL SKILLS',*FACT_MODELS]
    sections={}; active=None
    for line in lines:
        if line in headings: active=line; sections[active]=[]
        elif active: sections[active].append(line)
    profile.summary=' '.join(sections.get('PROFILE',[]))
    provenance='CV SHA256 '+hashlib.sha256(Path(path).read_bytes()).hexdigest()
    for line in sections.get('TECHNICAL SKILLS',[]):
        category,_,values=line.partition(':')
        for s in re.split(r'[,;]',values):
            if s.strip(): db.add(Skill(candidate_id=profile.id,text=s.strip(),context=category,provenance=provenance))
    # Keep source wording intact. A complete section is an evidence block, not a generated claim.
    for heading,model in FACT_MODELS.items():
        content='\n'.join(sections.get(heading,[]))
        if content: db.add(model(candidate_id=profile.id,text=content,provenance=provenance))
    if Path(path).resolve()!=(DATA/'master.pdf').resolve():
        handle,staged=tempfile.mkstemp(prefix='.master-',suffix='.pdf',dir=DATA);os.close(handle)
        try:
            shutil.copy2(path,staged);os.replace(staged,DATA/'master.pdf')
        finally: Path(staged).unlink(missing_ok=True)
    log(db,'Master CV imported; candidate confirmation required')
    db.flush(); return profile

def candidate(db):
    p=db.scalar(select(CandidateProfile))
    if not p: raise ValueError('Import the master CV first')
    return {**serialize(p),'skills':[serialize(x) for x in db.scalars(select(Skill))],**{m.__tablename__:[serialize(x) for x in db.scalars(select(m))] for m in FACT_MODELS.values()}}

def contains(text,word): return bool(re.search(r'(?<!\w)'+re.escape(word.lower())+r'(?!\w)',text.lower()))
def score(job,p,cfg,decision=None,item=None):
    """One authoritative scoring formula (issue #41): `decision['priority']`
    IS `decision['fit_assessment']['score']` (or evaluate_legacy()'s own
    priority in assessment_mode LEGACY) -- never a second, independently
    weighted total that gets silently discarded. `parts` (the legacy
    per-column UI breakdown) is a PROJECTION of the same FitAssessment
    components, never a second recomputation of fit.
    """
    from .recall import evaluate
    desc=job.description
    fields=('skills','employments','projects','education','certifications')
    evidence=(' '.join(str(f['text']) for key in fields for f in p.get(key,[])) if any(key in p for key in fields) else p.get('raw_text','')).lower()
    review=[]
    if re.search(r'UAE nationals?|Emirati|United Arab Emirates citizens?',job.title+' '+desc,re.I):
        review.append('Confirm UAE citizenship eligibility; UAE residence does not establish nationality')
    hard=[]
    for sentence in re.split(r'[\n.!?]',desc):
        if re.search(r'requir|must|mandatory|only|minimum',sentence,re.I): hard.append(sentence.strip())
    for h in hard:
        if re.search(r'national|citizen|clearance|sponsor|authoriz|visa',h,re.I): review.append('Confirm eligibility: '+h)
        if re.search(r'arabic',h,re.I): review.append('Confirm Arabic proficiency: '+h)
        if re.search(r'\bCCNA\b',h,re.I): review.append('Confirm the full CCNA certification; a course or keyword alone does not verify certification')
    if not desc.strip(): review.append('Job description is missing')

    item=item or {k:getattr(job,k,'') for k in ('title','description','location','remote_status','company','job_url','experience_requirement','date_posted','closing_date')}
    decision=decision or evaluate(item,cfg,p)
    fa=decision.get('fit_assessment')
    total=decision['priority'];blockers=[r['evidence'] for r in decision['hard']];review=list(dict.fromkeys(review+[r['evidence'] for r in decision['soft']]))
    matches=decision.get('matched_core_skills',[]);missing=decision.get('fit_assessment',{}).get('missing_skills',[])
    min_years=decision['experience']['minimum']
    if fa:
        c=fa['components']
        parts={'skills':c['profile_evidence']/20,'experience':c['experience']/15,'role':c['domain']/25,
               'education':1 if 'computer science' in evidence else .3,'location':c['geography']/10,
               'seniority':c['seniority']/10,'other':0 if review else 1}
        fit={'STRONG':'Strong fit','GOOD':'Good fit','STRETCH':'Stretch role','LOW':'Low priority — review','REJECTED':'Outside focus'}[fa['bucket']]
    else:
        local=any(contains(job.location,x) for x in cfg['locations'])
        role=any(norm(x) in norm(job.title) for x in cfg['target_roles'])
        parts={'skills':len(matches)/max(len(matches)+len(missing),1),'experience':1 if min_years==0 else (.6 if min_years==1 else .2),
               'role':1 if role else .2,'education':1 if 'computer science' in evidence else .3,'location':1 if local else .2,
               'seniority':0 if blockers else 1,'other':0 if review else 1}
        fit='Outside focus' if blockers else 'Stretch role' if min_years>=2 else 'Eligibility review' if review else 'Potential junior fit'
    if fa:
        # Issue #41 Phase 13: recommendation is a direct bucket projection,
        # never the legacy minimum_score/maybe_score thresholds -- those
        # don't align with the bucket boundaries and could otherwise map a
        # merely-LOW (not hard-rejected) job onto SKIP, which must be
        # reserved for an actual hard rejection (LOW stays NEEDS_REVIEW,
        # a normal rankable/visible disposition).
        recommendation='SKIP' if blockers else 'NEEDS_REVIEW' if review else {'STRONG':'HIGH_PRIORITY','GOOD':'APPLY','STRETCH':'MAYBE','LOW':'NEEDS_REVIEW','REJECTED':'SKIP'}[fa['bucket']]
    else:
        recommendation='SKIP' if blockers else 'NEEDS_REVIEW' if review else 'HIGH_PRIORITY' if total>=cfg['high_priority'] else 'APPLY' if total>=cfg['minimum_score'] else 'MAYBE' if total>=cfg['maybe_score'] else 'SKIP'
    score_kind=fa['score_kind'] if fa else 'Heuristic priority index — not a probability or eligibility decision. Never an interview, hiring, or success probability.'
    return dict(recall=decision,score=total,score_kind=score_kind,weights=cfg['weights'],fit=fit,minimum_experience_years=min_years,parts={k:round(v*100) for k,v in parts.items()},matching_skills=matches,missing_skills=missing,hard_requirements=hard,hard_blockers=blockers,needs_review=review,recommendation=recommendation,why=fit+'. Ranking priority, not a percentage match or hiring prediction. Skills not found are unknown, not proven absent. Experience, language and work authorization need your review. Edit profile facts and search preferences to correct the inputs.')

def _assessment_item(db,j):
    item={k:getattr(j,k,'') for k in ('title','description','location','remote_status','company','job_url',
                                      'experience_requirement','date_posted','closing_date')}
    observations=[]
    for o in db.scalars(select(JobObservation).where(JobObservation.job_id==j.id).order_by(JobObservation.id)):
        observations.append({k:getattr(o,k) for k in (
            'provider_family','provider_version','job_source_id','board_key','identity_kind','provider_job_id',
            'requisition_id','requisition_id_authority','employer_name_observed','title_observed',
            'location_observed','workplace_observed','description_observed','source_url','apply_url',
            'posted_at','posted_at_authority','closing_at','anomaly_flags','normalization_version')})
    item['_observation_authority']=observations
    return item


def preserve_workflow_state(db,j):
    """One policy for user-owned workflow state; assessment stays derived."""
    analysis=j.analysis or {}
    app=db.scalar(select(Application).where(Application.job_id==j.id))
    return bool(
        j.status in TERMINAL or app or
        analysis.get('saved') or analysis.get('bookmarked') or
        analysis.get('fit_at_application') or
        (app and app.tracking.get('fit_at_application')) or
        analysis.get('manual_classification') or analysis.get('workflow_status_preserved') or
        analysis.get('source_type')=='MANUAL' or str(j.source).casefold()=='manual'
    )


def analyze(db,j,profile=None,cfg=None):
    """Issue #41: `j.analysis['fit_assessment']` is the authoritative
    derived assessment (Owner Decision 3 -- no new table/column). A hard
    reject still gets a Job row and a JobObservation (#40 identity is
    untouched); it is only ever hidden through the existing SKIP status,
    never by skipping persistence, and existing saved/applied workflow
    state on this job is never overwritten here (see the status guard
    below, unchanged from before #41).
    """
    profile=profile if profile is not None else (candidate(db) if db.scalar(select(CandidateProfile)) else {})
    cfg=cfg or settings(db)
    item=_assessment_item(db,j)
    from .recall import evaluate
    decision=evaluate(item,cfg,profile)
    result=score(j,profile,cfg,decision=decision,item=item)
    fa=result['recall'].get('fit_assessment')
    j.analysis={**(j.analysis or {}),**result,
                **({'fit_assessment':fa} if fa else {}),
                **({'assessment_shadow':decision['assessment_shadow']} if decision.get('assessment_shadow') else {})}
    j.match_score=result['score']; j.matching_skills=result['matching_skills']; j.missing_skills=result['missing_skills']; j.red_flags=result['hard_blockers']+result['needs_review']; j.requirements=result['hard_requirements']; j.recommendation=result['recommendation']; j.priority='HIGH' if result['recommendation']=='HIGH_PRIORITY' else 'NORMAL'; j.hard_requirement_score=0 if j.red_flags else 100
    for field,key in [('skill_score','skills'),('experience_score','experience'),('education_score','education'),('location_score','location')]: setattr(j,field,result['parts'][key])
    # Owner Decision 1: a hard-rejected job is still hidden via the existing
    # SKIP status, but never by overwriting a job that already has saved/
    # application/terminal workflow state -- unchanged guard from before #41.
    if not preserve_workflow_state(db,j):
        j.status=result['recommendation'] if result['recommendation'] in STATUSES else 'ANALYZED'
    log(db,f'Analyzed: {j.match_score}/100, {j.recommendation}',j.id); return result

_PROVIDER_FAMILY_FOR_SOURCE={'Greenhouse':'greenhouse','Lever':'lever','Ashby':'ashby','SmartRecruiters':'smartrecruiters','LinkedIn':'linkedin'}

def _observation_from_record(record,data,job_source):
    """Rich provenance path (issue #40): `record` is the provider-native
    ProviderRecord backend.job_providers.compatibility.LegacyJobDict
    carried alongside the legacy ingestion dict, so Greenhouse/Lever/Ashby
    observations keep apply_url, identity kind, and provider facts that the
    legacy dict shape alone cannot express. Ashby is the only current
    provider whose identity can be a canonicalized URL fallback rather than
    a native id (issue #39 finding 6) -- raw_fields['identity_fallback']
    says so.
    """
    identity_kind='url_fallback' if record.raw_fields.get('identity_fallback')=='jobUrl' else 'native'
    return JobObservationInput(
        provider_family=record.provider,identity_kind=identity_kind,provider_job_id=record.provider_job_id,
        job_source_id=job_source.id if job_source else None,board_key=record.source_board,
        employer_name=data.get('company') or record.source_board or normalization.UNKNOWN,
        title=record.title,location=record.location,workplace=record.remote_status,description=record.description,
        source_url=record.source_url,apply_url=record.apply_url,posted_at=record.posted_at,
        posted_at_authority='documented_provider_field' if record.posted_at else 'none',
        closing_at=record.closing_at,retrieved_at=record.retrieved_at,provider_version=record.provider_version,
        provider_facts=dict(record.raw_fields),
    )

def _observation_from_data(data,job_source):
    """Fallback path for every discovery route with no rich ProviderRecord:
    SmartRecruiters (still legacy adapters.py transport -- issue #40 does
    not migrate it, only bridges its output into an observation), manual
    job entry, and tracker import. Never infers a fact `data` does not
    already carry.
    """
    source_label=data.get('source','Manual')
    provider_family=_PROVIDER_FAMILY_FOR_SOURCE.get(source_label,'manual')
    source_job_id=data.get('source_job_id','')
    identity_kind='native' if (provider_family!='manual' and source_job_id) else 'manual'
    posted_at=data.get('date_posted','')
    return JobObservationInput(
        provider_family=provider_family,identity_kind=identity_kind,provider_job_id=source_job_id,
        job_source_id=job_source.id if job_source else None,board_key=job_source.board if job_source else '',
        employer_name=data.get('company',normalization.UNKNOWN),title=data.get('title',''),
        location=data.get('location',normalization.UNKNOWN),workplace=data.get('remote_status',normalization.UNKNOWN),
        description=data.get('description',''),source_url=data.get('job_url',''),apply_url='',
        posted_at=posted_at,posted_at_authority='documented_provider_field' if (posted_at and provider_family!='manual') else 'none',
        closing_at=data.get('closing_date',''),retrieved_at=now(),provider_version='',provider_facts={},
    )

def apply_canonical_updates(job,observation):
    """Phase 9's canonical-field policy: a new observation may only FILL an
    UNKNOWN/empty canonical field, never overwrite an already-known value
    -- a field conflict stays evidence on the observation, never becomes an
    invented consensus on the Job.
    """
    view=canonical_view(observation)
    if job.company in ('',normalization.UNKNOWN) and view.company not in ('',normalization.UNKNOWN): job.company=view.company
    if job.location in ('',normalization.UNKNOWN) and view.location not in ('',normalization.UNKNOWN): job.location=view.location
    if job.remote_status in ('',normalization.UNKNOWN) and view.remote_status not in ('',normalization.UNKNOWN): job.remote_status=view.remote_status
    if not job.description and view.description: job.description=view.description
    if not job.date_posted and observation.posted_at_authority=='documented_provider_field' and view.date_posted: job.date_posted=view.date_posted
    if not job.closing_date and view.closing_date: job.closing_date=view.closing_date
    from .job_providers.contracts import valid_downstream_url
    stronger_source=(observation.provider_family not in ('manual','legacy') and
                     observation.identity_kind in ('native','url_fallback'))
    if not job.apply_url and view.apply_url and valid_downstream_url(view.apply_url): job.apply_url=view.apply_url
    if not job.job_url and stronger_source and valid_downstream_url(view.job_url):
        job.job_url=view.job_url
        job.canonical_url=view.canonical_url
    elif not job.canonical_url and job.job_url:
        job.canonical_url=canonical(job.job_url)

def _refresh_canonical_matching_keys(job):
    """Persist keys derived from canonical Job fields, never an observation."""
    job.normalized_employer_key=normalization.employer_key(job.company) or ''
    job.dedupe_fingerprint=deduplication.canonical_fingerprint(job)
    job.normalization_version=normalization.NORMALIZATION_VERSION

def _persist_observation(db,job,observation,decision):
    fingerprint=deduplication.composite_fingerprint(observation)
    ts=now()
    existing=None
    if observation.identity_kind in ('native','url_fallback') and observation.provider_job_id:
        existing=db.scalar(select(JobObservation).where(
            JobObservation.provider_family==observation.provider_family,
            JobObservation.job_source_id==observation.job_source_id,
            JobObservation.identity_kind==observation.identity_kind,
            JobObservation.provider_job_id==observation.provider_job_id))
    if existing is not None:
        existing.job_id=job.id; existing.employer_name_observed=observation.employer_name
        existing.title_observed=observation.title; existing.location_observed=observation.location
        existing.workplace_observed=observation.workplace; existing.description_observed=observation.description
        existing.source_url=observation.source_url; existing.apply_url=observation.apply_url
        existing.posted_at=observation.posted_at; existing.posted_at_authority=observation.posted_at_authority
        existing.closing_at=observation.closing_at; existing.retrieved_at=observation.retrieved_at or ts
        existing.provider_facts=dict(observation.provider_facts); existing.last_seen_at=ts
        existing.normalization_version=normalization.NORMALIZATION_VERSION
        existing.fingerprint=fingerprint; existing.match_method=decision.method
        existing.match_version=decision.version; existing.match_evidence=dict(decision.evidence)
        row=existing
    else:
        row=JobObservation(
            job_id=job.id,job_source_id=observation.job_source_id,provider_family=observation.provider_family,
            provider_version=observation.provider_version,board_key=observation.board_key,
            identity_kind=observation.identity_kind,provider_job_id=observation.provider_job_id,
            requisition_id=observation.requisition_id,requisition_id_authority=observation.requisition_id_authority,
            employer_name_observed=observation.employer_name,title_observed=observation.title,
            location_observed=observation.location,workplace_observed=observation.workplace,
            description_observed=observation.description,source_url=observation.source_url,
            apply_url=observation.apply_url,posted_at=observation.posted_at,
            posted_at_authority=observation.posted_at_authority,closing_at=observation.closing_at,
            retrieved_at=observation.retrieved_at or ts,provider_facts=dict(observation.provider_facts),
            anomaly_flags=list(observation.anomaly_flags),first_seen_at=ts,last_seen_at=ts,
            normalization_version=normalization.NORMALIZATION_VERSION,fingerprint=fingerprint,
            match_method=decision.method,match_version=decision.version,match_evidence=dict(decision.evidence))
        db.add(row)
    return row,fingerprint

def add_job(db,data,job_source=None):
    """Every discovery path -- Greenhouse/Lever/Ashby (rich ProviderRecord
    via backend.job_providers.compatibility.LegacyJobDict), SmartRecruiters
    (still its legacy adapters.py transport), manual job entry, and tracker
    import -- funnels through this one function (issue #40), so it is the
    single seam that builds a JobObservation and runs the conservative
    matcher (backend.deduplication) instead of the old O(n) all-Job
    fuzzy scan. `job_source` is the JobSource instance driving this
    discovery run, when there is one, so observation identity can be
    scoped to that source instance rather than provider family alone
    (backend.normalization: provider family != source instance).
    """
    provider_record=getattr(data,'provider_record',None)
    data={k:v for k,v in data.items() if k in Job.__table__.columns.keys() and k not in ('id','created_at','updated_at')}
    data.setdefault('company','UNKNOWN'); data.setdefault('title','Untitled role'); data.setdefault('location','UNKNOWN')
    for key in ('company','title','location','source','salary','remote_status'):
        if key in data and (not isinstance(data[key],str) or len(data[key])>500): raise ValueError('Job field is too long or has an invalid type')
    if len(data.get('description',''))>100000: raise ValueError('Job description exceeds 100,000 characters')
    if data.get('job_url'):
        from urllib.parse import urlsplit
        url=urlsplit(data['job_url'])
        if url.scheme not in ('http','https') or not url.hostname or url.username or url.password or len(data['job_url'])>2000: raise ValueError('Use a public HTTP(S) job link without credentials')
    if linkedin(data.get('job_url','')): data['source']='LinkedIn'
    data['canonical_url']=canonical(data.get('job_url',''))

    observation=_observation_from_record(provider_record,data,job_source) if provider_record is not None else _observation_from_data(data,job_source)
    decision=deduplication.resolve(db,observation)

    if decision.decision==deduplication.MATCH:
        j=db.get(Job,decision.selected_job_id)
        occurrences=list(j.analysis.get('occurrences',[]))
        for occurrence in [{'source':j.source,'url':j.job_url,'external_id':j.source_job_id},{'source':data.get('source','Manual'),'url':data.get('job_url',''),'external_id':data.get('source_job_id','')}]:
            if not any(o['source']==occurrence['source'] and o['url']==occurrence['url'] for o in occurrences): occurrences.append({**occurrence,'seen_at':now()})
        j.analysis={**j.analysis,'occurrences':occurrences}
        if data.get('notes') and data['notes'] not in j.notes: j.notes=(j.notes+'\n'+data['notes']).strip()
        if linkedin(j.job_url) and data.get('job_url') and not linkedin(data['job_url']): j.job_url=data['job_url']; j.canonical_url=data['canonical_url']; j.source=data.get('source','Company')
        apply_canonical_updates(j,observation)
        _persist_observation(db,j,observation,decision)
        _refresh_canonical_matching_keys(j)
        log(db,'Duplicate import merged into existing job',j.id)
        app=db.scalar(select(Application).where(Application.job_id==j.id))
        return j, {'duplicate_of':j.id,'application_id':app.id if app else None}

    data['normalized_employer_key']=normalization.employer_key(data.get('company')) or ''
    data['dedupe_fingerprint']=''; data['normalization_version']=normalization.NORMALIZATION_VERSION
    data.setdefault('apply_url',canonical_view(observation).apply_url)
    j=Job(**data); db.add(j); db.flush()
    _refresh_canonical_matching_keys(j)
    _persist_observation(db,j,observation,decision)
    log(db,'Job discovered',j.id); return j, None

def answer_for(db,q,p):
    if demographic(q) or live_assessment(q): return None
    a=db.scalar(select(ApprovedAnswer).where(ApprovedAnswer.normalized_question==norm(q),ApprovedAnswer.approved==True))
    if a: a.last_used=now(); return a.answer
    if sensitive(q): return None
    aliases={'full name':'name','name':'name','first name':'first_name','last name':'last_name','email':'email','email address':'email','phone':'phone','phone number':'phone','location':'location','current location':'location'}
    k=aliases.get(norm(q)); value=p.get(k) if k else None
    # Never split a person's full name to guess legally distinct name fields.
    return value if value and value!='UNKNOWN' else None

def generate_cv(db,j):
    p=candidate(db)
    if not p['confirmed']: raise ValueError('Confirm the extracted candidate profile before generating CVs')
    skills=sorted(p['skills'],key=lambda x:contains(j.description,x['text']),reverse=True)
    sections=[('Profile',p['summary']),('Technical Skills',', '.join(s['text'] for s in skills))]
    fact_ids=[f'skills:{s["id"]}' for s in skills]+[f'candidate_profiles:{p["id"]}:{field}' for field in ('name','location','email','phone','summary')]
    for title,key in [('Experience','employments'),('Projects','projects'),('Education','education'),('Certifications and Courses','certifications')]:
        for fact in p[key]: sections.append((title,fact['text'])); fact_ids.append(f'{key}:{fact["id"]}')
    tailored=p['name']+'\n'+p['location']+' | '+p['email']+' | '+p['phone']+'\n\n'+'\n\n'.join(h+'\n'+v for h,v in sections)
    stem=re.sub(r'[^\w-]+','_',f'Candidate_{j.id}')[:130]+f'_{datetime.now().strftime("%Y%m%d_%H%M%S_%f")}'
    directory=DATA/'documents'; directory.mkdir(exist_ok=True)
    pdf=directory/(stem+'_CV.pdf'); docx=directory/(stem+'_CV.docx')
    styles=getSampleStyleSheet(); styles['Normal'].fontSize=10; styles['Normal'].leading=14
    story=[Paragraph(escape(p['name']),styles['Title']),Paragraph(escape(p['location']+' | '+p['email']+' | '+p['phone']),styles['Normal'])]
    doc=Document(); doc.add_heading(p['name'],0); doc.add_paragraph(p['location']+' | '+p['email']+' | '+p['phone'])
    for h,v in sections:
        story.extend([Spacer(1,9),Paragraph(h,styles['Heading2'])]); doc.add_heading(h,1)
        for line in v.splitlines(): story.append(Paragraph(escape(line),styles['Normal'])); doc.add_paragraph(line)
    # Helvetica cannot render every writing system. Never silently ship missing glyphs.
    try: tailored.encode('cp1252'); pdf_supported=True
    except UnicodeEncodeError: pdf_supported=False
    if pdf_supported: SimpleDocTemplate(str(pdf),rightMargin=42,leftMargin=42,topMargin=36,bottomMargin=36).build(story)
    import unicodedata
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for paragraph in doc.paragraphs:
        if any(unicodedata.bidirectional(ch) in ('R','AL') for ch in paragraph.text):
            bidi=OxmlElement('w:bidi'); bidi.set(qn('w:val'),'1'); paragraph._p.get_or_add_pPr().append(bidi)
    doc.save(docx)
    row=ResumeVersion(job_id=j.id,pdf_path=str(pdf.relative_to(DATA)) if pdf_supported else '',docx_path=str(docx.relative_to(DATA)),master=p['raw_text'],tailored=tailored,diff='\n'.join(difflib.unified_diff(p['raw_text'].splitlines(),tailored.splitlines(),fromfile='MASTER',tofile='TAILORED',lineterm='')),fact_ids=fact_ids)
    db.add(row); db.flush(); log(db,'Tailored CV generated using source facts; skills reordered by relevance' if pdf_supported else 'Unicode DOCX generated. PDF unavailable for this script; review the DOCX in a suitable document editor.',j.id); return row
def cover(db,j):
    p=candidate(db)
    if not p['confirmed']: raise ValueError('Confirm the candidate profile before generating a cover letter')
    fact=p['summary']
    verified_skills=[s['text'] for s in p['skills'] if contains(j.description,s['text'])]
    text=f'Dear Hiring Team,\n\nI’m applying for the {j.title} position at {j.company}. {fact}\n\n'+('My background includes '+', '.join(verified_skills[:5])+'. ' if verified_skills else '')+f'I would welcome the opportunity to discuss how my background fits your team.\n\nKind regards,\n{p["name"]}'
    c=CoverLetter(job_id=j.id,text=text); db.add(c); log(db,'Cover letter generated for review',j.id); return c
def prepare(db,j):
    a=db.scalar(select(Application).where(Application.job_id==j.id))
    if a and (a.applied_date or a.status in TERMINAL or a.submission_attempted): raise ValueError('Application already submitted or awaiting submission reconciliation')
    analyze(db,j); resume=generate_cv(db,j); cfg=settings(db)
    if not a: a=Application(job_id=j.id); db.add(a); db.flush()
    if a.status in TERMINAL or a.submission_attempted: raise ValueError('Application already submitted or awaiting submission reconciliation')
    a.mode='PREPARE_ONLY'; a.status='NEEDS_REVIEW' if j.red_flags or j.recommendation=='SKIP' else 'READY_TO_APPLY'; j.status=a.status
    p=candidate(db)
    a.recruiter_message={'connection':f'Hello, I’m {p["name"]}. I’m interested in the {j.title} role at {j.company} and would appreciate connecting.', 'message':f'Hello, I’m interested in your {j.title} opening. Would you be open to discussing the role?', 'followup':f'Hello, I’m following up about the {j.title} opportunity at {j.company}. I remain interested and would welcome an update. Thank you, {p["name"]}.'}
    if cfg['cover_letter']=='always' or re.search(r'cover letter.{0,30}(required|mandatory)',j.description,re.I): cover(db,j)
    log(db,'Application prepared; LinkedIn requires manual submission' if linkedin(j.job_url) else 'Application prepared for review',j.id); return a
def set_status(db,j,status):
    transition(j.status,status); j.status=status
    a=db.scalar(select(Application).where(Application.job_id==j.id))
    if not a: a=Application(job_id=j.id); db.add(a); db.flush()
    transition(a.status,status); a.status=status
    if status=='APPLIED' and not a.applied_date:
        a.applied_date=now()
        if not a.confirmation: a.confirmation='USER REPORTED — receipt not independently verified'
        if not db.scalar(select(FollowUp).where(FollowUp.application_id==a.id)): db.add(FollowUp(application_id=a.id,due_date=(datetime.now(timezone.utc)+timedelta(days=settings(db)['followup_days'])).isoformat(),message=f'Following up on my application for {j.title} at {j.company}. I remain interested and would appreciate an update.'))
    from .campaign import record_event
    if a.tracking.get('stage')!=status:
        a.tracking={**a.tracking,'stage':status}
        record_event(db,a,status,source='USER')
    log(db,f'Status changed to {status}',j.id); return a

HEADERS={'Date Found':'date_found','Company':'company','Job Title':'title','Location':'location','Work Mode':'remote_status','Job URL':'job_url','Source':'source','Job Description / Notes from JD':'description','Match Score':'match_score','Priority':'priority','Missing Skills':'missing_skills','Status':'status','Salary':'salary','Notes':'notes'}
def import_tracker(db,path):
    from .document_security import validate_document
    validate_document(Path(path).read_bytes(),Path(path).name,'xlsx')
    w=load_workbook(path,data_only=True,read_only=True,keep_links=False); count=0; dup=0
    ws=w['Applications'] if 'Applications' in w.sheetnames else w.active
    if (ws.max_column or 0)>100 or (ws.max_row or 0)>10001:
        w.close(); raise ValueError('Tracker dimensions exceed 100 columns or 10,000 data rows')
    from itertools import islice
    rows=list(islice(ws.iter_rows(max_col=100,values_only=True),10002))
    if not rows or len(rows)>10001: w.close(); raise ValueError('Tracker must contain 1 to 10,000 data rows')
    if sum(len(str(value)) for row in rows for value in row if value is not None)>2_000_000:
        w.close(); raise ValueError('Tracker text exceeds 2 million characters')
    header=[str(x or '').strip() for x in rows[0]]
    if 'Company' not in header or 'Job Title' not in header: raise ValueError('Tracker needs Company and Job Title columns')
    for row in rows[1:]:
        raw=dict(zip(header,row))
        if not raw.get('Company') or not raw.get('Job Title'): continue
        data={v:str(raw[k]) for k,v in HEADERS.items() if raw.get(k) is not None}
        data['match_score']=int(float(raw.get('Match Score') or 0)); data['missing_skills']=str(raw.get('Missing Skills') or '').split(',') if raw.get('Missing Skills') else []
        state=norm(str(raw.get('Status') or 'FOUND')).upper().replace(' ','_'); data['status']=state if state in STATUSES else 'FOUND'
        j,d=add_job(db,data); dup+=bool(d); count+=not bool(d)
        j.analysis={**(j.analysis or {}),'source_type':'TRACKER','workflow_status_preserved':True}
        if data['status'] in TERMINAL and not d:
            a=Application(job_id=j.id,status=data['status'],applied_date=str(raw.get('Applied Date') or '')); db.add(a); db.flush()
            if raw.get('Recruiter'): db.add(Recruiter(application_id=a.id,name=str(raw['Recruiter']),contacted=str(raw.get('Recruiter Contacted','')).lower() in ('yes','true')))
            if raw.get('Follow-up Date'): db.add(FollowUp(application_id=a.id,due_date=str(raw['Follow-up Date'])))
            if raw.get('Interview Date'): db.add(Interview(application_id=a.id,date=str(raw['Interview Date'])))
        analyze(db,j)
    w.close(); log(db,f'Tracker imported: {count} records, {dup} duplicates'); return {'imported':count,'duplicates':dup,'sheets':w.sheetnames}
def safe_cell(value): return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value
def sync_tracker(db):
    if settings(db).get('campaign_workbook'):
        from .workbook import export_workbook
        return export_workbook(db,DATA/'tracker.xlsx')
    target=DATA/'tracker.xlsx'; backup=DATA/'backups'; backup.mkdir(exist_ok=True)
    if target.exists(): shutil.copy2(target,backup/f'tracker_{datetime.now().strftime("%Y%m%d_%H%M%S_%f")}.xlsx'); w=load_workbook(target)
    else:
        w=Workbook(); w.active.title='Applications'; w.active.append(['ID',*HEADERS,'CV Version','Applied Date','Recruiter','Recruiter Contacted','Follow-up Date','Interview Date'])
    ws=w['Applications']; cols={str(c.value):c.column for c in ws[1] if c.value}
    existing={str(ws.cell(r,cols.get('Job URL',7)).value):r for r in range(2,ws.max_row+1) if ws.cell(r,cols.get('Company',3)).value}
    for j in db.scalars(select(Job)):
        r=existing.get(j.job_url) if j.job_url else next((i for i in range(2,ws.max_row+1) if cols.get('ID') and str(ws.cell(i,cols['ID']).value)==str(j.id)),None)
        if not r: r=next((i for i in range(2,ws.max_row+2) if not ws.cell(i,cols['Company']).value),ws.max_row+1)
        a=db.scalar(select(Application).where(Application.job_id==j.id)); cv=db.scalar(select(ResumeVersion).where(ResumeVersion.job_id==j.id).order_by(ResumeVersion.id.desc()))
        vals={k:getattr(j,v) for k,v in HEADERS.items()}; vals.update({'ID':j.id,'CV Version':Path(cv.pdf_path).name if cv else '', 'Applied Date':a.applied_date if a else ''})
        if a:
            f=db.scalar(select(FollowUp).where(FollowUp.application_id==a.id)); recruiter=db.scalar(select(Recruiter).where(Recruiter.application_id==a.id)); interview=db.scalar(select(Interview).where(Interview.application_id==a.id).order_by(Interview.date.desc()))
            vals.update({'Follow-up Date':f.due_date if f else '', 'Recruiter':recruiter.name if recruiter else '', 'Recruiter Contacted':'Yes' if recruiter and recruiter.contacted else 'No','Interview Date':interview.date if interview else ''})
        for key,val in vals.items():
            if key not in cols: continue
            c=ws.cell(r,cols[key])
            if not c.has_style: c._style=copy(ws.cell(2,cols[key])._style)
            if isinstance(val,list): val=', '.join(val)
            if key in ('Date Found','Applied Date','Follow-up Date','Interview Date') and val:
                try: val=datetime.fromisoformat(str(val)).replace(tzinfo=None); c.number_format='dd mmm yyyy'
                except ValueError: pass
            c.value=safe_cell(val)
    tmp=target.with_suffix('.tmp.xlsx'); w.save(tmp); w.close(); check=load_workbook(tmp); check.close(); os.replace(tmp,target)
    log(db,'Excel synchronized; timestamped backup preserved'); return {'path':'tracker.xlsx','rows':db.query(Job).count()}
