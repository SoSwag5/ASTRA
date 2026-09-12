"""Database-derived workbook. Atomic output and durable retry state."""
import os,shutil,threading,math
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.chart import BarChart,Reference
from openpyxl.utils import get_column_letter
from .models import Session,WorkbookSync,DATA,settings,now
from .campaign import application_rows,analytics,campaign_settings
from .search_workspace import parse_date

sync_lock=threading.Lock()
HEADERS=['Application ID','Company','Job Title','Location','Emirate','Source','Job URL','Date Found','Date Applied','Fit','Status','Current Stage','CV Version','Salary','Last Activity','Next Action','Next Action Date','Follow-up Date','Outcome','Rejection Reason','Notes']

def safe(value):
    return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value

def sheet(book,name,headers,rows,table_name):
    if name in book:book.remove(book[name])
    ws=book.create_sheet(name);ws.append(headers)
    for row in rows:ws.append([safe(v) for v in row])
    ws.freeze_panes='C2';ws.sheet_view.showGridLines=False
    for cell in ws[1]:cell.font=Font(name='Aptos',bold=True,color='FFFFFF',size=11);cell.fill=PatternFill('solid',fgColor='183A46');cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws.row_dimensions[1].height=32
    for row in ws.iter_rows(min_row=2):
        for c in row:c.font=Font(name='Aptos',size=11);c.alignment=Alignment(vertical='top',wrap_text=True)
    for col in ws.columns:
        header=str(col[0].value);width=45 if header in ('Notes','Job URL','Next Action') else 30 if header in ('Company','Job Title','CV Version') else 19
        ws.column_dimensions[col[0].column_letter].width=width
        if 'Date' in header or header=='Last Activity':
            for cell in col[1:]:
                date=parse_date(cell.value)
                if date:cell.value=date.replace(tzinfo=None);cell.number_format='dd mmm yyyy'
    for row in ws.iter_rows(min_row=2):
        lines=max((sum(max(1,math.ceil(len(part)/max(8,ws.column_dimensions[c.column_letter].width-3))) for part in str(c.value or '').split('\n')) for c in row),default=1)
        ws.row_dimensions[row[0].row].height=min(409,max(24,lines*16+8))
    if rows:
        table=Table(displayName=table_name,ref=f'A1:{ws.cell(ws.max_row,len(headers)).coordinate}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
    return ws

def export_workbook(db,target):
    target=Path(target);target.parent.mkdir(parents=True,exist_ok=True)
    backup=target.parent/'backups';backup.mkdir(exist_ok=True)
    if target.exists():
        shutil.copy2(target,backup/f'tracker_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx')
        book=load_workbook(target,keep_links=False)
    else:book=Workbook();book.remove(book.active)
    # Retain legacy job inventory once, plus all unrelated user sheets.
    if 'Applications' in book and book['Applications']['A1'].value!='Application ID':
        book['Applications'].title='Discovered Jobs'
    if 'How to Use' in book:
        guide=book['How to Use']
        for merged in list(guide.merged_cells.ranges):
            if merged.min_row<=2:guide.unmerge_cells(str(merged))
        guide.merge_cells('A1:C2');guide['A1']='Campaign upgrade: edit application records in the app. Applications and Activity are regenerated from the database; Excel edits are not imported automatically. The older instructions below describe the retained legacy workbook.'
        guide['A1'].alignment=Alignment(wrap_text=True,vertical='center');guide['A1'].font=Font(name='Aptos',size=11,bold=True,color='FFFFFF');guide['A1'].fill=PatternFill('solid',fgColor='183A46');guide.row_dimensions[1].height=30;guide.row_dimensions[2].height=30
    rows=application_rows(db);stats=analytics(rows,campaign_settings(db)['observation_days'])
    values=[];events=[]
    for r in rows:
        j=r['job'];t=r['tracking'];f=r['followup']
        values.append([r['id'],j['company'],j['title'],j['location'],r['emirate'],t.get('source_at_application') or j['source'],j['job_url'],j['date_found'],r['applied_date'],t.get('fit_at_application','Not recorded'),r['status'],r['stage'],t.get('cv_version','Not recorded'),j['salary'],t.get('last_activity') or r['updated_at'],t.get('next_action',''),t.get('next_action_date',''),f['due_date'] if f and not f['done'] else '',t.get('outcome',''),t.get('rejection_reason',''),j['notes']])
        for e in r['events']:events.append([e['occurred_at'],r['id'],j['company'],j['title'],e['event_type'],e['stage'],e['message'],e['source']])
    sheet(book,'Applications',HEADERS,values,'CampaignApplications')
    sheet(book,'Activity',['Date','Application ID','Company','Role','Event','Stage','Notes','Origin'],sorted(events,key=lambda e:e[0]),'CampaignActivity')
    metrics=[['Submitted applications',stats['total']],['Active applications',stats['active']],['Applications this week',stats['this_week']],['Applications this month',stats['this_month']],['Meaningful responses',stats['responses']],['Interviews',stats['interviews']],['Offers',stats['offers']],['Mature response rate',stats['response_rate'] if stats['response_rate'] is not None else 'Not enough data'],['Interview rate',stats['interview_rate'] if stats['interview_rate'] is not None else 'Not enough data'],['Offer rate',stats['offer_rate'] if stats['offer_rate'] is not None else 'Not enough data']]
    data=sheet(book,'Dashboard Data',['Metric','Value'],metrics,'CampaignMetrics')
    if 'Dashboard' in book:book.remove(book['Dashboard'])
    dash=book.create_sheet('Dashboard',0);dash.sheet_view.showGridLines=False
    dash.merge_cells('A1:L2');dash['A1']='UAE job search';dash['A1'].font=Font(name='Aptos Display',size=26,bold=True,color='183A46')
    dash.merge_cells('A3:L3');dash['A3']='Local application record · Updated '+datetime.now().strftime('%d %b %Y %H:%M')
    for index,(label,value) in enumerate(metrics):
        row=5+(index//5)*4;col=1+(index%5)*3
        dash.cell(row,col,label).font=Font(name='Aptos',size=11,color='516570')
        cell=dash.cell(row+1,col,f"='Dashboard Data'!B{index+2}");cell.font=Font(name='Aptos',size=16 if isinstance(value,str) else 22,bold=True,color='183A46')
        dash.row_dimensions[row+1].height=34
        cell.alignment=Alignment(vertical='center',horizontal='left')
        dash.merge_cells(start_row=row+1,start_column=col,end_row=row+1,end_column=col+2)
        if 'rate' in label:cell.number_format='0%';data.cell(index+2,2).number_format='0.0%'
    for col in range(1,16):dash.column_dimensions[get_column_letter(col)].width=12
    dash.merge_cells('A13:O14');dash['A13']=f"Response rate: meaningful, non-automated replies among applications at least {stats['observation_days']} days old (n={stats['mature_count']}). Interview and offer rates use all submitted applications. Historical CV and fit stay unrecorded when unknown."
    dash['A13'].alignment=Alignment(wrap_text=True,vertical='top');dash['A13'].font=Font(name='Aptos',size=11)
    for offset,(title,labels,counts) in enumerate([
        ('Applications by week',[r['week'] for r in stats['weeks']],[r['applications'] for r in stats['weeks']]),
        ('Current pipeline',list(stats['pipeline']),list(stats['pipeline'].values())),
        ('Applications by source',[r['name'] for r in stats['groups']['source']],[r['applications'] for r in stats['groups']['source']]),
        ('Applications by emirate',[r['name'] for r in stats['groups']['emirate']],[r['applications'] for r in stats['groups']['emirate']])]):
        if not counts:continue
        col=4+offset*3;data.cell(1,col,title);data.cell(1,col+1,'Applications')
        data.column_dimensions[get_column_letter(col)].width=30;data.column_dimensions[get_column_letter(col+1)].width=15
        for cell in [data.cell(1,col),data.cell(1,col+1)]:cell.font=Font(name='Aptos',bold=True,color='FFFFFF',size=11);cell.fill=PatternFill('solid',fgColor='183A46');cell.alignment=Alignment(wrap_text=True)
        for i,(label,count) in enumerate(zip(labels,counts),2):data.cell(i,col,safe(label));data.cell(i,col+1,count)
        chart=BarChart();chart.title=title;chart.style=10;chart.height=8;chart.width=18
        chart.add_data(Reference(data,min_col=col+1,min_row=1,max_row=len(counts)+1),titles_from_data=True);chart.set_categories(Reference(data,min_col=col,min_row=2,max_row=len(counts)+1));chart.legend=None;chart.y_axis.majorUnit=1;chart.y_axis.numFmt='0'
        dash.add_chart(chart,('A17','I17','A34','I34')[offset])
    temporary=target.with_name(target.stem+'.tmp.xlsx')
    try:
        book.save(temporary);book.close()
        check=load_workbook(temporary,read_only=True)
        ids=[r[0] for r in check['Applications'].iter_rows(min_row=2,values_only=True) if r[0] is not None]
        check.close()
        if len(ids)!=len(rows) or len(set(ids))!=len(ids):raise ValueError('Workbook application validation failed')
        os.replace(temporary,target)
        for old in sorted(backup.glob('tracker_????????_??????_??????.xlsx'),reverse=True)[30:]:
            if old.is_file() and not old.is_symlink():old.unlink()
    finally:
        book.close()
        if temporary.exists():temporary.unlink()
    return {'path':target.name,'rows':len(rows),'events':len(events)}

def retry_sync():
    if not sync_lock.acquire(False):return
    from .main import task_lock
    if not task_lock.acquire(False):sync_lock.release();return
    try:
        with Session() as db:
            state=db.get(WorkbookSync,1)
            if not state or state.revision<=state.exported_revision or not settings(db).get('campaign_workbook'):return
            revision=state.revision
            try:
                export_workbook(db,DATA/'tracker.xlsx')
                state.exported_revision=revision;state.last_success=now();state.error=''
            except Exception:
                state.error='Excel update pending. Close the workbook if open; the app will retry. Your application data is saved.'
            db.commit()
    finally:task_lock.release();sync_lock.release()
